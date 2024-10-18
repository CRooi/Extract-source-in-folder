import os
from docx import Document


def extract_code_to_word(folder_path, output_file):
    # 创建一个Word文档对象
    doc = Document()
    doc.add_heading('Extracted Code Files', level=1)

    # 遍历指定文件夹
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            # 检查文件扩展名，以只提取代码文件
           if file.endswith(('.java','.vue')):  # 可根据需要添加更多扩展名
            # if file.endswith(('.py', '.java', '.c', '.cpp', '.js', '.html', '.css', '.rb', '.go')):  # 可根据需要添加更多扩展名
                file_path = os.path.join(root, file)
                doc.add_heading(file, level=2)

                # 读取文件内容并添加到Word文档中
                with open(file_path, 'r', encoding='utf-8') as code_file:
                    content = code_file.read()
                    doc.add_paragraph(content)
                    doc.add_paragraph()  # 添加一个空行以分隔文件内容

    # 保存Word文档
    doc.save(output_file)
    print(f"代码已成功提取并保存到 {output_file}")


# 使用示例
folder_path = r'项目文件夹'  # 替换为代码文件夹的路径，可以包含前后端代码

output_file = 'qianduan_code.docx'  # 输出Word文件名

extract_code_to_word(folder_path, output_file)
